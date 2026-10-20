import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="D3D3D3", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'<w:tblBorders {nsdecls("w")}><w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/><w:left w:val="none"/><w:right w:val="none"/></w:tblBorders>')
    tblPr.append(borders)

def make_doc():
    doc = Document()
    
    # Page setup - Margins: Top=2cm, Bottom=2cm, Left=2.5cm, Right=2cm
    for section in doc.sections:
        section.top_margin = Inches(0.79)
        section.bottom_margin = Inches(0.79)
        section.left_margin = Inches(0.98)
        section.right_margin = Inches(0.79)
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)

    # Styles setup
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Times New Roman'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = p_title.add_run("BÀI TẬP LỚN KỸ THUẬT LẬP TRÌNH - ĐỀ TÀI 7: QUẢN LÝ DỊCH VỤ TRUYỀN HÌNH SỐ / IPTV\n")
    run_sub.font.name = 'Times New Roman'
    run_sub.font.size = Pt(12)
    run_sub.font.bold = True
    run_sub.font.color.rgb = RGBColor(0x00, 0x33, 0x66)

    run_main = p_title.add_run("PHỤ LỤC B. BẢNG ĐẶC TẢ VÀ TEST CASE USE CASE\n")
    run_main.font.name = 'Times New Roman'
    run_main.font.size = Pt(15)
    run_main.font.bold = True
    run_main.font.color.rgb = RGBColor(0x99, 0x00, 0x00)

    run_author = p_title.add_run("Sinh viên thực hiện: Dương Bảo Long | Phụ trách: Package (Gói dịch vụ) & SupportTicket (Phiếu hỗ trợ kỹ thuật)")
    run_author.font.name = 'Times New Roman'
    run_author.font.size = Pt(11)
    run_author.font.italic = True
    run_author.font.color.rgb = RGBColor(0x44, 0x44, 0x44)

    doc.add_paragraph() # Spacing

    # =========================================================================
    # PART 1: USE CASE 1 - PACKAGE
    # =========================================================================
    h1 = doc.add_paragraph()
    r = h1.add_run("I. ĐẶC TẢ USE CASE 1: QUẢN LÝ GÓI KÊNH / DỊCH VỤ (PACKAGE)")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12.5)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x00, 0x33, 0x66)

    # Specification Table Package
    table_pkg_spec = doc.add_table(rows=0, cols=2)
    table_pkg_spec.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table_pkg_spec)

    pkg_spec_data = [
        ("Mã use case", "UC06 (hoặc phân hệ Gói cước / Package trong hệ thống TV36)"),
        ("Tên use case", "Quản lý Gói dịch vụ truyền hình (Package Management)"),
        ("Sinh viên thực hiện", "Họ tên: Dương Bảo Long                     MSSV: ...................."),
        ("Lớp đối tượng (C++)", "Package (Entities/Package.h, Entities/Package.cpp)\nLớp giao diện điều khiển: PackageMenu (UseCases/PackageMenu.h, UseCases/PackageMenu.cpp)"),
        ("Lớp cơ sở kế thừa", "Entity (Entities/Entity.h) - Lớp trừu tượng định nghĩa định danh id và các hàm thuần ảo: parseFromString, exportToString, exportToVector."),
        ("Danh sách thuộc tính và ràng buộc", 
         "1. id (string): Kế thừa từ Entity. Định dạng 'GOI' + số tự tăng (ví dụ GOI001), sinh tự động bởi Repository thông qua packages_counter.txt, duy nhất, không thể sửa đổi.\n"
         "2. name (string): Tên gói dịch vụ. Không được để trống, không chứa ký tự '|' hoặc ký tự xuống dòng (\\r, \\n).\n"
         "3. monthlyPrice (long long): Cước phí thuê bao tháng (VNĐ). Ràng buộc phải là số nguyên dương (> 0).\n"
         "4. status (enum PackageStatus): Trạng thái gói, gồm: DANG_CUNG_CAP (Đang cung cấp) và NGUNG_DANG_KY (Ngừng đăng ký).\n"
         "5. channelIDs (vector<string>): Danh sách mã kênh thuộc gói. Phải có nhiều hơn 1 kênh (> 1 kênh), không được trùng lặp, các mã kênh phải hợp lệ (bắt đầu 'KENH') và phải tồn tại trong Data/channels.txt (Channel Repository).\n"
         "6. vodIDs (vector<string>): Danh sách mã nội dung VOD kèm theo. Phải có nhiều hơn 1 VOD (> 1 VOD), không được trùng lặp, các mã VOD phải hợp lệ (bắt đầu 'VOD') và phải tồn tại trong Data/noidungvod.txt (NoiDungVOD Repository)."
        ),
        ("Tên file dữ liệu", "Data/packages.txt (Lưu trữ danh sách bản ghi gói cước)\nData/packages_counter.txt (Lưu trữ bộ đếm số tự tăng sinh mã tự động)"),
        ("Định dạng một bản ghi", "id|name|monthlyPrice|status|channelIDs|vodIDs\nVí dụ: GOI001|Goi Gia Dinh|120000|NGUNG_DANG_KY|KENH001,KENH002,KENH003|VOD001,VOD002"),
        ("Tham chiếu tới (khóa ngoại)", 
         "- Channel (Data/channels.txt): Kiểm tra toàn bộ mã trong channelIDs phải tồn tại thực tế trong kho kênh.\n"
         "- NoiDungVOD (Data/noidungvod.txt): Kiểm tra toàn bộ mã trong vodIDs phải tồn tại thực tế trong kho nội dung VOD."
        ),
        ("Được tham chiếu bởi", 
         "- Contract (Data/contracts.txt): Hợp đồng thuê bao tham chiếu tới Package thông qua packageID. Khi xóa hoặc cập nhật gói, hệ thống kiểm tra ràng buộc toàn vẹn đối với các hợp đồng đang hoạt động (ACTIVE)."
        ),
        ("Chức năng Thêm (C)", 
         "- Tự động sinh ID duy nhất (GOIxxx) từ packages_counter.txt.\n"
         "- Nhập tên gói (kiểm tra không rỗng, không chứa ký tự phân cách '|').\n"
         "- Nhập cước phí tháng dạng số thực/số nguyên (kiểm tra > 0).\n"
         "- Chọn trạng thái gói qua menu số trực quan: 1. DANG_CUNG_CAP, 2. NGUNG_DANG_KY.\n"
         "- Nhập danh sách kênh phân tách bằng dấu phẩy: kiểm tra định dạng 'KENHxxx', kiểm tra tồn tại từng kênh trong repoChannel, chống trùng lặp, bắt buộc số lượng > 1 kênh.\n"
         "- Nhập danh sách VOD phân tách bằng dấu phẩy: kiểm tra định dạng 'VODxxx', kiểm tra tồn tại từng VOD trong repoVOD, chống trùng lặp, bắt buộc số lượng > 1 VOD.\n"
         "- Cơ chế bắt lỗi: Mọi giá trị nhập đều đi qua setter với try-catch và ghi log qua Logger/Debugger. Thêm vào bộ nhớ và đồng bộ ngay lập tức ra file packages.txt."
        ),
        ("Chức năng Xem, tìm kiếm (R)", 
         "- Xem tất cả: Xuất danh sách toàn bộ các gói dạng bảng căn cột chuẩn hoá bằng TablePrinter (các cột: ID, Name, Monthly Price, Status, Channels, VODs).\n"
         "- Tìm kiếm theo mã ID: Nhập mã gói (tự động chuẩn hoá chữ in hoa), tìm kiếm chính xác qua Repository::readByID(). Nếu tìm thấy hiển thị chi tiết dạng bảng đơn (TablePrinter::printSingle); nếu không thông báo lỗi."
        ),
        ("Chức năng Cập nhật (U)", 
         "- Nhập mã gói cần sửa, kiểm tra tồn tại.\n"
         "- Kiểm tra ràng buộc hợp đồng đang hoạt động (ACTIVE) qua isPackageInActiveContract():\n"
         "  + Nếu gói đang nằm trong Hợp đồng có trạng thái ACTIVE: Hệ thống khoá sửa đổi thông tin nghiệp vụ, CHỈ cho phép cập nhật Trạng thái (Status).\n"
         "  + Nếu gói không nằm trong Hợp đồng ACTIVE: Cho phép sửa từng trường (Tên, Giá cước, Trạng thái, Danh sách Kênh, Danh sách VOD); ấn Enter (bỏ trống) để giữ nguyên giá trị cũ.\n"
         "- Toàn bộ dữ liệu sửa đổi đều được kiểm tra tính hợp lệ và toàn vẹn tham chiếu trước khi ghi lại file."
        ),
        ("Chức năng Xóa (D)", 
         "- Nhập mã gói cần xóa, hiển thị chi tiết bản ghi trước khi thực hiện.\n"
         "- Kiểm tra ràng buộc toàn vẹn tham chiếu (Removal dependency constraint): Nếu gói đang được gán cho bất kỳ Hợp đồng nào có trạng thái ACTIVE (c.get_packageID() == packageID && c.get_status() == ContractStatus::ACTIVE), hệ thống CHẶN không cho xóa và đưa ra thông báo cảnh báo rõ ràng.\n"
         "- Nếu không vi phạm ràng buộc: Yêu cầu người dùng xác nhận chắc chắn (y/n). Sau khi xác nhận 'y', thực hiện xóa khỏi Repository và cập nhật ngay vào file packages.txt."
        ),
        ("Kỹ thuật OOP áp dụng", 
         "- Đóng gói (Encapsulation): Các trường name, monthlyPrice, status, channelIDs, vodIDs đều để private; truy xuất qua getter/setter, trong các setter ném ngoại lệ (throw invalid_argument) khi vi phạm quy tắc dữ liệu.\n"
         "- Kế thừa & Trừu tượng (Inheritance & Abstraction): Kế thừa lớp trừu tượng Entity, ghi đè (override) 3 phương thức thuần ảo: parseFromString(), exportToString(), exportToVector().\n"
         "- Xử lý ngoại lệ (Exception Handling): Sử dụng try-catch-throw kiểm soát dữ liệu đầu vào và chuyển đổi chuỗi (stoll, splitString).\n"
         "- Lớp khuôn mẫu (Template & Generic Programming): Sử dụng Repository<Package> để tái sử dụng toàn bộ logic quản lý CRUD và thao tác I/O file bền vững.\n"
         "- Kiểu liệt kê strongly typed (enum class): Sử dụng enum class PackageStatus đảm bảo an toàn kiểu dữ liệu."
        )
    ]

    for label, val in pkg_spec_data:
        row = table_pkg_spec.add_row()
        c1, c2 = row.cells[0], row.cells[1]
        c1.width = Inches(1.8)
        c2.width = Inches(4.7)
        set_cell_background(c1, "F2F4F7")
        set_cell_margins(c1, 80, 80, 100, 100)
        set_cell_margins(c2, 80, 80, 100, 100)
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(label)
        r1.font.name = 'Times New Roman'
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
        
        p2 = c2.paragraphs[0]
        r2 = p2.add_run(val)
        r2.font.name = 'Times New Roman'
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    doc.add_paragraph() # Spacing

    # Test cases table Package
    p_tc_title = doc.add_paragraph()
    r = p_tc_title.add_run("Bảng Test Case kiểm thử Use Case Quản lý Gói kênh (Package):")
    r.font.name = 'Times New Roman'
    r.font.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(0x00, 0x33, 0x66)

    table_pkg_tc = doc.add_table(rows=0, cols=6)
    table_pkg_tc.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table_pkg_tc)

    tc_headers = ["STT", "Chức năng", "Dữ liệu vào", "Kết quả mong đợi", "Kết quả thực tế", "Đánh giá"]
    hdr_row = table_pkg_tc.add_row()
    widths = [Inches(0.4), Inches(1.1), Inches(1.8), Inches(1.8), Inches(1.0), Inches(0.4)]
    for i, h in enumerate(tc_headers):
        cell = hdr_row.cells[i]
        cell.width = widths[i]
        set_cell_background(cell, "003366")
        set_cell_margins(cell, 100, 100, 80, 80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.font.name = 'Times New Roman'
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    pkg_tc_data = [
        ("1", "Thêm gói", "Name: Gói Siêu Cấp\nPrice: -50000", "Báo lỗi [Validation Error] Giá cước tháng phải lớn hơn 0 và yêu cầu nhập lại", "Chương trình thông báo lỗi giá cước, không lưu, cho nhập lại", "Đạt"),
        ("2", "Thêm gói", "Name: Gói Phim | Hay (chứa ký tự '|')\nPrice: 150000", "Báo lỗi tên gói không được chứa ký tự '|' hoặc xuống dòng", "Báo lỗi hợp lệ, ngăn chặn phá vỡ định dạng file", "Đạt"),
        ("3", "Thêm gói", "Channel IDs: KENH999 (mã kênh không tồn tại trong channels.txt)", "Báo lỗi Channel ID 'KENH999' does not exist in repository và yêu cầu nhập lại", "Hiển thị lỗi và yêu cầu nhập lại toàn bộ danh sách kênh", "Đạt"),
        ("4", "Thêm gói", "Channel IDs: KENH001 (chỉ nhập 1 kênh)", "Báo lỗi Package must contain more than 1 channel (> 1 channel)", "Bắt buộc gói phải có từ 2 kênh trở lên, yêu cầu nhập lại", "Đạt"),
        ("5", "Thêm gói", "VOD IDs: VOD001,VOD001 (trùng lặp mã VOD)", "Báo lỗi Duplicate VOD ID 'VOD001' in list", "Phát hiện mã trùng lặp, yêu cầu nhập lại danh sách VOD", "Đạt"),
        ("6", "Thêm gói", "Dữ liệu hợp lệ: Tên 'Goi 4K VIP', Giá 190000, Status 1, Channels: KENH001,KENH002, VODs: VOD001,VOD002", "Tạo thành công gói mới với mã sinh tự động (ví dụ GOI011), lưu vào packages.txt", "Tạo thành công, sinh mã tự động, dữ liệu ghi nhận ngay vào file", "Đạt"),
        ("7", "Xem gói", "Tìm kiếm theo mã: goi002 (nhập chữ thường)", "Chuẩn hoá thành 'GOI002', hiển thị đúng bảng chi tiết gói Cơ bản", "Chuẩn hoá chữ in hoa và in ra thông tin chính xác", "Đạt"),
        ("8", "Xem gói", "Tìm kiếm theo mã: GOI999 (không tồn tại)", "Thông báo lỗi [Error] Package not found, ID: GOI999", "Hiển thị thông báo không tìm thấy, không crash chương trình", "Đạt"),
        ("9", "Cập nhật", "Gói GOI001 (đang nằm trong Hợp đồng HD001 có trạng thái ACTIVE)", "Thông báo cảnh báo gói đang trong hợp đồng ACTIVE, khoá sửa tên/giá/kênh/VOD, chỉ cho sửa Status", "Chỉ hiển thị tuỳ chọn cập nhật trạng thái gói, bảo vệ hợp đồng", "Đạt"),
        ("10", "Cập nhật", "Gói chưa gắn hợp đồng ACTIVE: Sửa Tên mới, các trường khác ấn Enter bỏ qua", "Cập nhật tên mới, các thuộc tính còn lại giữ nguyên giá trị ban đầu", "Cập nhật đúng tên mới, dữ liệu khác được bảo toàn nguyên vẹn", "Đạt"),
        ("11", "Xóa gói", "Gói GOI002 (đang gắn vào hợp đồng ACTIVE HD002)", "Từ chối xóa, thông báo 'Cannot delete Package because it is currently assigned to an ACTIVE contract!'", "Chặn thao tác xóa thành công, bảo toàn tính toàn vẹn dữ liệu", "Đạt"),
        ("12", "Xóa gói", "Gói không có hợp đồng ACTIVE: Chọn 'n' khi hỏi xác nhận (y/n)", "Huỷ thao tác xóa, giữ nguyên bản ghi trong file và bộ nhớ", "Thông báo Deletion cancelled, dữ liệu được giữ nguyên", "Đạt"),
        ("13", "Xóa gói", "Gói không có hợp đồng ACTIVE: Chọn 'y' khi hỏi xác nhận", "Xóa thành công, cập nhật file packages.txt, không còn xuất hiện khi tìm kiếm", "Bản ghi bị xóa thành công khỏi file", "Đạt")
    ]

    for row_data in pkg_tc_data:
        row = table_pkg_tc.add_row()
        for i, val in enumerate(row_data):
            cell = row.cells[i]
            cell.width = widths[i]
            set_cell_margins(cell, 60, 60, 60, 60)
            p = cell.paragraphs[0]
            if i in [0, 4, 5]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(val)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(8.5)
            if i == 5:
                r.font.bold = True
                r.font.color.rgb = RGBColor(0x00, 0x66, 0x00)

    doc.add_page_break()

    # =========================================================================
    # PART 2: USE CASE 2 - SUPPORT TICKET
    # =========================================================================
    h2 = doc.add_paragraph()
    r = h2.add_run("II. ĐẶC TẢ USE CASE 2: QUẢN LÝ PHIẾU HỖ TRỢ KỸ THUẬT (SUPPORT TICKET)")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12.5)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0x00, 0x33, 0x66)

    # Specification Table Support Ticket
    table_st_spec = doc.add_table(rows=0, cols=2)
    table_st_spec.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table_st_spec)

    st_spec_data = [
        ("Mã use case", "UC10 (hoặc phân hệ Phiếu hỗ trợ kỹ thuật trong hệ thống TV36)"),
        ("Tên use case", "Quản lý Phiếu hỗ trợ kỹ thuật (Support Ticket Management)"),
        ("Sinh viên thực hiện", "Họ tên: Dương Bảo Long                     MSSV: ...................."),
        ("Lớp đối tượng (C++)", "SupportTicket (Entities/SupportTicket.h, Entities/SupportTicket.cpp)\nLớp giao diện điều khiển: SupportTicketMenu (UseCases/SupportTicketMenu.h, UseCases/SupportTicketMenu.cpp)"),
        ("Lớp cơ sở kế thừa", "Entity (Entities/Entity.h) - Lớp cơ sở trừu tượng chứa mã id và các hàm thuần ảo: parseFromString, exportToString, exportToVector."),
        ("Danh sách thuộc tính và ràng buộc", 
         "1. id (string): Kế thừa từ Entity. Mã phiếu dạng 'PHT' + số tự tăng (ví dụ PHT001), sinh tự động bởi Repository thông qua support_tickets_counter.txt, duy nhất, không sửa đổi.\n"
         "2. contractID (string): Mã hợp đồng phát sinh sự cố. Không được để trống, đúng định dạng 'HDxxx' và bắt buộc phải tồn tại trong Data/contracts.txt.\n"
         "3. stbID (string): Mã thiết bị Set-Top Box liên quan. Tuỳ chọn (để trống nếu lỗi không do thiết bị như cước phí, đổi gói). Nếu nhập, phải bắt đầu bằng 'STB' và phải khớp đúng với thiết bị đang gán cho Hợp đồng đó.\n"
         "4. receivedDate (string): Thời điểm tiếp nhận yêu cầu (định dạng 'yyyy-mm-dd HH:mm'). Tự động lấy timestamp hiện tại của hệ thống, không được để trống, không chứa ký tự '|'.\n"
         "5. content (string): Mô tả nội dung sự cố, khiếu nại của khách hàng. Không được để trống, không chứa ký tự '|' hoặc ký tự xuống dòng.\n"
         "6. priority (enum TicketPriority): Mức độ ưu tiên, gồm: BINH_THUONG, UU_TIEN, KHAN_CAP.\n"
         "7. resolution (string): Kết quả / phương án xử lý sự cố. Mặc định ban đầu để trống khi mới tiếp nhận, không chứa ký tự '|'.\n"
         "8. closedDate (string): Thời điểm đóng phiếu hỗ trợ (định dạng 'yyyy-mm-dd HH:mm'). Trống khi đang xử lý; tự động suy diễn trạng thái: nếu đã có resolution hoặc closedDate thì trạng thái là DA_DONG (Đã đóng), ngược lại là TIEP_NHAN (Tiếp nhận)."
        ),
        ("Tên file dữ liệu", "Data/support_tickets.txt (Lưu trữ các phiếu hỗ trợ kỹ thuật)\nData/support_tickets_counter.txt (Lưu trữ bộ đếm số tự tăng sinh mã phiếu)"),
        ("Định dạng một bản ghi", "id|contractID|stbID|receivedDate|content|priority|resolution|closedDate\nVí dụ: PHT001|HD001|STB001|2026-09-20 08:00|STB khong len nguon|BINH_THUONG|Da kiem tra va thay the bo xu ly|2026-09-20 10:00"),
        ("Tham chiếu tới (khóa ngoại)", 
         "- Contract (Data/contracts.txt): Hợp đồng tham chiếu bắt buộc; kiểm tra sự tồn tại của hợp đồng trước khi tạo phiếu.\n"
         "- SetTopBox (Data/settopboxes.txt): Thiết bị liên quan trong hợp đồng; kiểm tra tính hợp lệ và sự đồng bộ với thiết bị trong hợp đồng."
        ),
        ("Được tham chiếu bởi", 
         "- Customer / Contract: Là lịch sử hỗ trợ kỹ thuật và chăm sóc khách hàng của từng hợp đồng/khách hàng.\n"
         "- Ràng buộc nghiệp vụ đặc biệt: Dữ liệu phiếu hỗ trợ kỹ thuật mang tính chất nhật ký kiểm toán (Audit Trail / CSKH), do đó KHÔNG ĐƯỢC PHÉP XÓA để bảo toàn lịch sử phục vụ."
        ),
        ("Chức năng Thêm (C)", 
         "- Tự động sinh ID duy nhất (PHTxxx) từ counter.\n"
         "- Nhập mã hợp đồng: kiểm tra cú pháp 'HDxxx', nạp và tra cứu từ repoContract; nếu không tồn tại yêu cầu nhập lại.\n"
         "- Tra cứu thiết bị STB: Hệ thống tự động đọc stbID gắn với hợp đồng đó; cho phép người dùng nhấn Enter để lấy tự động STB của hợp đồng hoặc bỏ trống nếu không liên quan tới thiết bị; chặn nếu nhập STB không khớp với hợp đồng.\n"
         "- Tự động sinh thời điểm tiếp nhận (receivedDate) theo giờ hệ thống thực tế chuẩn 'yyyy-mm-dd HH:mm'.\n"
         "- Nhập nội dung sự cố (kiểm tra không rỗng, chống ký tự '|').\n"
         "- Chọn mức độ ưu tiên qua menu số: 1. BINH_THUONG, 2. UU_TIEN, 3. KHAN_CAP.\n"
         "- Mặc định khởi tạo resolution và closedDate là rỗng (phiếu ở trạng thái TIEP_NHAN).\n"
         "- Lưu vào Repository và đồng bộ ngay vào file support_tickets.txt."
        ),
        ("Chức năng Xem, tìm kiếm (R)", 
         "- Xem tất cả: Xuất danh sách toàn bộ phiếu hỗ trợ kỹ thuật dạng bảng căn cột đẹp mắt qua TablePrinter (các cột: ID, Contract, STB ID, Received Date, Content, Priority, Resolution, Closed Date).\n"
         "- Tìm kiếm theo mã ID: Nhập mã phiếu PHTxxx (chuẩn hoá chữ hoa), tìm kiếm chính xác qua Repository::readByID(). Nếu tìm thấy hiển thị chi tiết qua TablePrinter::printSingle; nếu không thông báo lỗi."
        ),
        ("Chức năng Cập nhật (U)", 
         "- Nhập mã phiếu cần cập nhật, kiểm tra tồn tại.\n"
         "- Ràng buộc phụ thuộc tham chiếu (Dependency Rule): Để bảo toàn tính pháp lý và lịch sử gốc, hệ thống KHÔNG CHO PHÉP sửa ID, Contract ID, STB ID và Thời điểm tiếp nhận (Received Date).\n"
         "- CHỈ cho phép cập nhật 4 trường nghiệp vụ giải quyết sự cố:\n"
         "  1. Content (Nội dung bổ sung / sửa đổi).\n"
         "  2. Priority (Mức độ ưu tiên mới).\n"
         "  3. Solution / Resolution (Phương án, kết quả xử lý của kỹ thuật viên).\n"
         "  4. Close Date (Thời điểm hoàn thành đóng phiếu, cho phép gõ 'now' để tự điền timestamp hiện tại).\n"
         "- Cho phép ấn Enter để giữ nguyên giá trị cũ từng trường; kiểm tra hợp lệ và lưu file."
        ),
        ("Chức năng Xóa (D)", 
         "- Nhập mã phiếu cần xóa, hiển thị chi tiết thông tin phiếu.\n"
         "- Ràng buộc nghiệp vụ kiểm toán (Removal dependency constraint):\n"
         "  Hệ thống hiển thị thông báo lỗi từ chối xóa: '[Error] Cannot delete Support Ticket ... [Rule] Removal dependency constraint: Support tickets cannot be deleted (must be kept for audit history)'.\n"
         "- Thiết kế này hoàn toàn tuân thủ tiêu chí thiết kế hệ thống thực tế trong viễn thông: dữ liệu khiếu nại, sự cố kỹ thuật không bao giờ được phép xóa để phục vụ tra soát, hậu kiểm và KPI kỹ thuật."
        ),
        ("Kỹ thuật OOP áp dụng", 
         "- Đóng gói (Encapsulation): Toàn bộ thuộc tính nghiệp vụ để private; bảo vệ dữ liệu bằng các phương thức getter/setter kèm logic validate chống dữ liệu bẩn và ném exception.\n"
         "- Kế thừa & Trừu tượng (Inheritance & Abstraction): Kế thừa lớp trừu tượng Entity, ghi đè 3 phương thức ảo thuần tuý parseFromString, exportToString, exportToVector.\n"
         "- Phương thức dẫn xuất / logic suy diễn: Phương thức getStatus() và getStatusString() tự động tính toán dựa trên việc có closedDate/resolution hay chưa, vừa tương thích các module khác (như SetTopBoxMenu) vừa tinh gọn dữ liệu lưu trữ.\n"
         "- Generic Repository: Sử dụng Repository<SupportTicket> tái sử dụng mã nguồn lưu trữ và thao tác bảng.\n"
         "- Quản lý bộ nhớ và xử lý lỗi: Áp dụng bắt lỗi toàn diện qua try-catch, phân tách file .h/.cpp chuẩn chỉnh theo chuẩn C++17."
        )
    ]

    for label, val in st_spec_data:
        row = table_st_spec.add_row()
        c1, c2 = row.cells[0], row.cells[1]
        c1.width = Inches(1.8)
        c2.width = Inches(4.7)
        set_cell_background(c1, "F2F4F7")
        set_cell_margins(c1, 80, 80, 100, 100)
        set_cell_margins(c2, 80, 80, 100, 100)
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(label)
        r1.font.name = 'Times New Roman'
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = RGBColor(0x11, 0x11, 0x11)
        
        p2 = c2.paragraphs[0]
        r2 = p2.add_run(val)
        r2.font.name = 'Times New Roman'
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = RGBColor(0x22, 0x22, 0x22)

    doc.add_paragraph() # Spacing

    # Test cases table Support Ticket
    p_tc_title2 = doc.add_paragraph()
    r = p_tc_title2.add_run("Bảng Test Case kiểm thử Use Case Quản lý Phiếu hỗ trợ kỹ thuật (Support Ticket):")
    r.font.name = 'Times New Roman'
    r.font.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(0x00, 0x33, 0x66)

    table_st_tc = doc.add_table(rows=0, cols=6)
    table_st_tc.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table_st_tc)

    hdr_row2 = table_st_tc.add_row()
    for i, h in enumerate(tc_headers):
        cell = hdr_row2.cells[i]
        cell.width = widths[i]
        set_cell_background(cell, "003366")
        set_cell_margins(cell, 100, 100, 80, 80)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(h)
        r.font.name = 'Times New Roman'
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    st_tc_data = [
        ("1", "Thêm phiếu", "Contract ID: HD999 (hợp đồng không tồn tại)", "Báo lỗi [Error] Contract ID 'HD999' does not exist in repository! và yêu cầu nhập lại", "Báo lỗi hợp lệ, ngăn chặn tạo phiếu cho hợp đồng ảo", "Đạt"),
        ("2", "Thêm phiếu", "Contract ID: HD001 (đang gán STB001), nhưng nhập STB ID: STB002", "Báo lỗi STB ID must match the STB assigned to this Contract (STB001)!", "Chặn nhập sai thiết bị của hợp đồng, yêu cầu nhập lại", "Đạt"),
        ("3", "Thêm phiếu", "STB ID: Nhấn phím Enter (bỏ trống)", "Tự động gán mã STB từ hợp đồng HD001 (STB001) vào phiếu", "Tự động lấy STB001 từ hợp đồng, tối ưu trải nghiệm nhập", "Đạt"),
        ("4", "Thêm phiếu", "Content: Để trống (ấn Enter ngay)", "Báo lỗi không được để trống trường nội dung sự cố", "Báo lỗi hợp lệ, yêu cầu nhập nội dung cụ thể", "Đạt"),
        ("5", "Thêm phiếu", "Dữ liệu hợp lệ: Contract HD001, Content 'Khong vao duoc kenh k+', Priority 3 (KHAN_CAP)", "Tạo thành công phiếu mới với mã tự sinh PHTxxx, tự điền Received Date hiện tại, Resolution và Closed Date để trống", "Tạo thành công, sinh mã tự động, lưu tức thì vào file support_tickets.txt", "Đạt"),
        ("6", "Xem phiếu", "Tìm kiếm theo mã: pht001 (nhập chữ thường)", "Chuẩn hoá thành 'PHT001', in bảng thông tin chi tiết của phiếu PHT001", "Chuẩn hoá in hoa và hiển thị đầy đủ thông tin", "Đạt"),
        ("7", "Xem phiếu", "Tìm kiếm theo mã: PHT999 (không tồn tại)", "Thông báo lỗi [Error] Ticket not found, ID: PHT999", "Báo không tìm thấy phiếu, không crash chương trình", "Đạt"),
        ("8", "Cập nhật", "Phiếu PHT002: Cập nhật Solution 'Da reset lai port tin hieu', Close Date gõ 'now'", "Cập nhật kết quả xử lý và tự động lấy timestamp hiện tại cho Close Date; trạng thái phiếu tự động chuyển thành DA_DONG", "Cập nhật thành công, file ghi nhận ngày giờ hiện tại", "Đạt"),
        ("9", "Cập nhật", "Phiếu PHT001: Enter bỏ qua toàn bộ các trường", "Không thay đổi bất kỳ trường dữ liệu nào, bản ghi giữ nguyên vẹn", "Bản ghi giữ nguyên nội dung cũ", "Đạt"),
        ("10", "Xóa phiếu", "Chọn Xóa phiếu PHT001 (hoặc bất kỳ phiếu nào)", "Hệ thống từ chối xóa, thông báo lỗi: 'Support tickets cannot be deleted (must be kept for audit history)'", "Chặn xóa tuyệt đối nhằm đảm bảo toàn vẹn nhật ký hỗ trợ khách hàng", "Đạt")
    ]

    for row_data in st_tc_data:
        row = table_st_tc.add_row()
        for i, val in enumerate(row_data):
            cell = row.cells[i]
            cell.width = widths[i]
            set_cell_margins(cell, 60, 60, 60, 60)
            p = cell.paragraphs[0]
            if i in [0, 4, 5]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(val)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(8.5)
            if i == 5:
                r.font.bold = True
                r.font.color.rgb = RGBColor(0x00, 0x66, 0x00)

    # Save document
    output_path = r"C:\Users\default.LAPTOP-EV5NLL8H\Desktop\IPTV\TV36\PhuLucB_DacTa_Package_SupportTicket_DuongBaoLong.docx"
    doc.save(output_path)
    print(f"File created successfully at: {output_path}")

if __name__ == "__main__":
    make_doc()
